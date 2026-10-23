from io import BytesIO
import re
from recommendation_display import display_fields
from recommendations import category_for
from datetime import date,datetime
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

def safe(value):return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value

def product_type(row):
 ai_type=row.get('ai_product_type')
 if isinstance(ai_type,str) and ai_type.strip(): return ai_type.strip()
 title=row.get('title')
 text=(title or '').lower()
 if re.search(r'\b(?:replacement|protective|charging) case\b',text) and re.search(r'\bearbuds?\b|\btws\b|\bairpods\b',text): return 'Earbud Case'
 types=[
  ('Electric Toothbrush',r'\b(?:electric|sonic|power)\s+tooth\s?brush(?:es)?\b'),
  ('Toothbrush',r'\btooth\s?brush(?:es)?\b'),
  ('Earbuds',r'\bear\s?buds?\b|\btws\b|\b(?:nord|galaxy|oneplus) buds\b|\bairpods\b'),
  ('Headphones',r'\bheadphones?\b|\bheadsets?\b'),
  ('Earphones',r'\bearphones?\b|\bin[- ]ear\b'),
  ('Soundbar',r'\bsound\s?bars?\b'),
  ('Speaker',r'\bspeakers?\b'),
  ('Smartwatch',r'\bsmart\s?watches?\b|\bfitness (?:band|tracker)\b'),
  ('Smartphone',r'\bsmartphones?\b|\bmobile phones?\b|\biphone\b'),
  ('Laptop',r'\blaptops?\b|\bmacbooks?\b|\bchromebooks?\b'),
  ('Tablet',r'\btablets?\b|\bipad\b'),
  ('Television',r'\btelevisions?\b|\bsmart tv\b|\b(?:led|oled|qled) tv\b'),
  ('Monitor',r'\b(?:computer|gaming|desktop) monitors?\b'),
  ('Keyboard',r'\bkeyboards?\b'),
  ('Mouse',r'\b(?:wireless|gaming|optical|computer) mouse\b'),
  ('Air Cooler',r'\bair coolers?\b'),
  ('Air Conditioner',r'\bair conditioners?\b|\b(?:split|window|inverter) ac\b'),
  ('Refrigerator',r'\brefrigerators?\b|\bfridges?\b'),
  ('Washing Machine',r'\bwashing machines?\b'),
  ('Vacuum Cleaner',r'\bvacuum cleaners?\b'),
  ('Water Purifier',r'\bwater purifiers?\b'),
  ('Mixer Grinder',r'\bmixer grinders?\b'),
  ('Electric Kettle',r'\b(?:electric )?kettles?\b'),
  ('Air Fryer',r'\bair fryers?\b'),
  ('Hair Dryer',r'\bhair dryers?\b'),
  ('Trimmer',r'\btrimmers?\b'),
  ('Backpack',r'\bbackpacks?\b'),
  ('Shoes',r'\bshoes?\b|\bsneakers?\b'),
 ]
 for label,pattern in types:
  if re.search(pattern,text): return label
 category=category_for(title) if title else row.get('category','general')
 labels={'fan':'Fan','air_purifier':'Air Purifier','heater':'Room Heater','humidifier':'Humidifier','dehumidifier':'Dehumidifier','power_bank':'Power Bank / Charger','water_bottle':'Water Bottle','cleaning':'Cleaning Product','kitchen':'Kitchen Product','textiles':'Textiles','sports':'Sports Product'}
 return labels.get(category,'Unknown / Other')

def export_report(job,results,recommendations):
 wb=Workbook();comparison=wb.active;comparison.title='City Comparison';advice=wb.create_sheet('Recommendations')
 matrix=len(job['config']['cities'])>1
 headers=['Brand','ASIN'] if matrix else ['Brand','ASIN','Product Title','City','Pincode','Days','Delivery Date','Status','Price (INR)','Error / Notes']
 if matrix:
  for city in job['config']['cities']:
   headers.extend([label+' ('+city['name']+')' for label in ('Days','Delivery Date','Status','Price')])
 headers.extend(['Product Type','Past Month Bought'])
 comparison.append(headers)
 lookup={(row['asin'],row['pincode']):row for row in results};metadata={}
 for row in results:
  if row.get('title'):metadata[row['asin']]=row
 ai_types={row['asin']:row['ai_product_type'] for row in recommendations if row.get('ai_product_type')}
 for product in job['config']['products']:
  info=dict(metadata.get(product['asin'],{}))
  if product['asin'] in ai_types: info['ai_product_type']=ai_types[product['asin']]
  if matrix:
   values=[safe(info.get('brand') or product.get('brand')),product['asin']]
   for city in job['config']['cities']:
    row=lookup.get((product['asin'],city['pincode']),{})
    values.extend([row.get('delivery_days'),date.fromisoformat(row['delivery_date']) if row.get('delivery_date') else None,row.get('status','Not checked'),row.get('price')])
   values.extend([product_type(info),safe(info.get('purchase_badge') or 'Unavailable')])
   comparison.append(values)
   comparison.cell(comparison.max_row,2).comment=Comment('Product: '+str(info.get('title') or '')+'\nSource: https://www.amazon.in/dp/'+product['asin'],'Delivery Atlas')
  for city in job['config']['cities']:
   row=lookup.get((product['asin'],city['pincode']),{})
   if not matrix: comparison.append([safe(info.get('brand') or product.get('brand')),product['asin'],safe(info.get('title')),city['name'],city['pincode'],row.get('delivery_days'),date.fromisoformat(row['delivery_date']) if row.get('delivery_date') else None,row.get('status','Not checked'),row.get('price'),safe(row.get('error')),product_type(info),safe(row.get('purchase_badge') or 'Unavailable')])
   status_column=5+4*job['config']['cities'].index(city) if matrix else 8
   comparison.cell(comparison.max_row,status_column).comment=Comment('City: '+city['name']+'\nPincode: '+city['pincode']+'\nNotes: '+str(row.get('error') or '')+'\nPromise: '+str(row.get('delivery_text') or 'Unavailable')+'\nPincode verified: '+str(row.get('location_verified',False))+'\nChecked: '+str(row.get('checked_at','Not checked'))+'\nSource: '+row.get('source_url','https://www.amazon.in/dp/'+product['asin']),'Delivery Atlas')
 advice.append(['ASIN','Product Name','Product Type','General Recommendation','Seasonal Recommendation','Past Month Bought'])
 for row in sorted(recommendations,key=lambda r:(r['asin'],r['city'])):
  row=dict(row)
  row.setdefault('ai_status',job['config'].get('ai_status') or ('AI assessment pending; run is still in progress' if job.get('status')=='running' else 'AI assessment was not recorded for this run; run a new check'))
  text=display_fields(row)
  advice.append([row['asin'],safe(row.get('title')),product_type(row),safe(text['general_recommendation']),safe(text['seasonal_recommendation']),safe(row.get('purchase_badge') or 'Unavailable')])
  advice.cell(advice.max_row,5).comment=Comment('Engine: '+str(row.get('ai_engine','Local rules'))+'\n'+str(row.get('basis','')),'Delivery Atlas')
 widths={comparison:([18,16]+[14,23,24,18]*len(job['config']['cities'])+[24,34] if matrix else [16,15,36,17,11,8,16,21,14,35,24,34]),advice:[15,40,24,60,85,34]}
 for sheet in wb:
  sheet.freeze_panes=('C2' if matrix else 'D2') if sheet==comparison else 'D2';sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.zoomScale=85;sheet.row_dimensions[1].height=35
  sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows='1:1'
  for index,width in enumerate(widths[sheet],1):sheet.column_dimensions[get_column_letter(index)].width=width
  for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='244B7E' if sheet==comparison and matrix else '173F46');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(vertical='center',wrap_text=True)
  for row in sheet.iter_rows(min_row=2):
   for cell in row:
    cell.alignment=Alignment(vertical='top',wrap_text=True)
    if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F0F6F5')
    if isinstance(cell.value,date):cell.number_format='dd-mmm-yyyy'
   lines=max(len(str(cell.value or ''))//max(10,widths[sheet][cell.column-1]-3)+1 for cell in row);sheet.row_dimensions[row[0].row].height=22 if sheet==comparison and matrix else max(55,min(220,lines*15))
  if sheet==comparison:
   for row in sheet.iter_rows(min_row=2):
    if matrix:
     for offset in range(2,len(row)-2,4):
      days,delivery,status,price=row[offset:offset+4]
      if isinstance(days.value,(int,float)):
       days.fill=PatternFill('solid',fgColor='C6EFCE' if days.value<=2 else 'D9E2F3' if days.value<=5 else 'FFEB9C')
      price.number_format='#,##0.00'
      status.fill=PatternFill('solid',fgColor='C6EFCE' if status.value=='Available' else 'FFF2CC' if status.value=='Not checked' else 'FFC7CE');status.font=Font(bold=True)
      delivery.number_format='dd-mm-yyyy'
    else:
     row[8].number_format='#,##0.00'
     row[7].fill=PatternFill('solid',fgColor='C6EFCE' if row[7].value=='Available' else 'FFC7CE');row[7].font=Font(bold=True)
 stream=BytesIO();wb.save(stream);stream.seek(0);return stream
